
import csv
import io
import json
import logging
import mimetypes
import os
import tempfile
import zipfile
from collections.abc import Generator
from dataclasses import dataclass
from types import SimpleNamespace
from typing import Any

import requests
from dify_plugin import Tool
from dify_plugin.entities.tool import ToolInvokeMessage
from dify_plugin.errors.tool import ToolProviderCredentialValidationError
from PIL import Image, ImageOps
from yarl import URL

logger = logging.getLogger(__name__)

SUPPORTED_MINERU_EXTENSIONS = {".pdf", ".doc", ".docx", ".ppt", ".pptx", ".png", ".jpg", ".jpeg"}
SUPPORTED_FALLBACK_EXTENSIONS = {".md", ".markdown", ".html", ".htm", ".json", ".yaml", ".yml", ".txt"}
SUPPORTED_EXTENSIONS = SUPPORTED_MINERU_EXTENSIONS | SUPPORTED_FALLBACK_EXTENSIONS
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg"}
OFFICE_ZIP_EXTENSIONS = {".docx", ".pptx"}
PDF_EXTENSION = ".pdf"


@dataclass
class Credentials:
    base_url: str
    server_type: str
    token: str | None = None


@dataclass
class NormalizedFile:
    filename: str
    blob: bytes
    report: dict[str, Any]


class MmbMaterialPreprocessorTool(Tool):
    def _invoke(self, tool_parameters: dict[str, Any]) -> Generator[ToolInvokeMessage, None, None]:
        credentials = self._get_credentials()
        source_file = tool_parameters.get("file")
        if source_file is None:
            raise ValueError("file is required")

        original_name = getattr(source_file, "filename", "document") or "document"
        extension = os.path.splitext(original_name)[1].lower()
        if extension not in SUPPORTED_EXTENSIONS:
            raise ValueError(
                f"Unsupported file extension {extension}. Supported extensions: "
                f"{sorted(SUPPORTED_EXTENSIONS)}"
            )

        normalized = self._normalize_file(source_file, tool_parameters)
        file_metadata = self._build_file_metadata(source_file, normalized, extension)

        yield self.create_variable_message("file_metadata", file_metadata)
        yield self.create_variable_message("compression_report", normalized.report)

        if extension in SUPPORTED_MINERU_EXTENSIONS:
            md_content, content_list = self._parse_local(credentials, normalized, tool_parameters)
        else:
            md_content, content_list = self._parse_fallback_document(source_file, extension)
        md_content = self._with_metadata_header(md_content, file_metadata, normalized.report)

        yield self.create_variable_message("content_list", content_list)
        yield self.create_variable_message("content_list_text", json.dumps(content_list, ensure_ascii=False))
        yield self.create_text_message(md_content)
        yield self.create_json_message(
            {
                "file_metadata": file_metadata,
                "compression_report": normalized.report,
                "content_list": content_list,
            }
        )

    def _get_credentials(self) -> Credentials:
        base_url = (self.runtime.credentials.get("base_url") or "").rstrip("/")
        server_type = self.runtime.credentials.get("server_type", "local")
        token = self.runtime.credentials.get("token")
        if not base_url:
            raise ToolProviderCredentialValidationError("Please input MinerU base_url")
        if server_type != "local":
            raise ToolProviderCredentialValidationError("MMB material preprocessor currently supports local MinerU only")
        return Credentials(base_url=base_url, server_type=server_type, token=token)

    @staticmethod
    def _headers(credentials: Credentials) -> dict[str, str]:
        if credentials.token:
            return {"Authorization": f"Bearer {credentials.token}", "accept": "application/json"}
        return {"accept": "application/json"}

    @staticmethod
    def _api_url(base_url: str, *paths: str) -> str:
        return str(URL(base_url) / "/".join(paths))

    def _normalize_file(self, file: Any, tool_parameters: dict[str, Any]) -> NormalizedFile:
        filename = getattr(file, "filename", "document") or "document"
        blob = getattr(file, "blob", b"") or b""
        extension = os.path.splitext(filename)[1].lower()
        report: dict[str, Any] = {
            "strategy": "balanced",
            "original_filename": filename,
            "normalized_filename": filename,
            "original_size_bytes": len(blob),
            "normalized_size_bytes": len(blob),
            "compressed": False,
            "actions": [],
            "warnings": [],
        }

        max_long_edge = self._number(tool_parameters.get("max_long_edge"), 1800)
        jpeg_quality = self._number(tool_parameters.get("jpeg_quality"), 82)
        large_file_mb = self._number(tool_parameters.get("large_file_mb"), 30)
        large_threshold = int(large_file_mb * 1024 * 1024)

        try:
            if extension in IMAGE_EXTENSIONS:
                filename, blob, actions = self._compress_image_file(filename, blob, max_long_edge, jpeg_quality)
                report["actions"].extend(actions)
            elif extension in OFFICE_ZIP_EXTENSIONS and len(blob) >= large_threshold:
                blob, actions = self._compress_office_media(blob, extension, max_long_edge, jpeg_quality)
                report["actions"].extend(actions)
            elif extension == PDF_EXTENSION and len(blob) >= large_threshold:
                pdf_blob, actions, warning = self._compress_scanned_pdf(blob, max_long_edge, jpeg_quality)
                if pdf_blob is not None:
                    blob = pdf_blob
                    report["actions"].extend(actions)
                elif warning:
                    report["warnings"].append(warning)
            elif len(blob) < large_threshold:
                report["actions"].append("skipped: below large_file_mb threshold")
        except Exception as exc:
            logger.warning("Failed to normalize %s: %s", filename, exc)
            report["warnings"].append(f"normalization failed; original file used: {exc}")

        report["normalized_filename"] = filename
        report["normalized_size_bytes"] = len(blob)
        report["compressed"] = len(blob) < report["original_size_bytes"]
        if not report["actions"]:
            report["actions"].append("skipped: format preserved for parsing quality")
        return NormalizedFile(filename=filename, blob=blob, report=report)

    @staticmethod
    def _number(value: Any, default: int) -> int:
        try:
            return int(value)
        except (TypeError, ValueError):
            return default

    def _compress_image_file(
        self, filename: str, blob: bytes, max_long_edge: int, jpeg_quality: int
    ) -> tuple[str, bytes, list[str]]:
        with Image.open(io.BytesIO(blob)) as image:
            image = ImageOps.exif_transpose(image)
            original_size = image.size
            image = self._resize_image(image, max_long_edge)
            out = io.BytesIO()
            if image.mode in {"RGBA", "LA", "P"}:
                background = Image.new("RGB", image.convert("RGBA").size, (255, 255, 255))
                background.paste(image.convert("RGBA"), mask=image.convert("RGBA").split()[-1])
                image = background
            else:
                image = image.convert("RGB")
            stem = os.path.splitext(filename)[0] or "image"
            image.save(out, format="JPEG", quality=jpeg_quality, optimize=True, progressive=True)
            return f"{stem}.jpg", out.getvalue(), [f"image resized {original_size} -> {image.size}, jpeg quality {jpeg_quality}"]

    def _compress_office_media(
        self, blob: bytes, extension: str, max_long_edge: int, jpeg_quality: int
    ) -> tuple[bytes, list[str]]:
        media_prefix = "word/media/" if extension == ".docx" else "ppt/media/"
        actions: list[str] = []
        out = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(blob), "r") as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                lower_name = item.filename.lower()
                if lower_name.startswith(media_prefix) and lower_name.endswith((".jpg", ".jpeg", ".png")):
                    try:
                        data, action = self._compress_embedded_image(data, lower_name, max_long_edge, jpeg_quality)
                        actions.append(f"{item.filename}: {action}")
                    except Exception as exc:
                        actions.append(f"{item.filename}: skipped ({exc})")
                zout.writestr(item, data)
        if not actions:
            actions.append("office media compression checked: no compressible media found")
        return out.getvalue(), actions

    def _compress_embedded_image(
        self, data: bytes, name: str, max_long_edge: int, jpeg_quality: int
    ) -> tuple[bytes, str]:
        with Image.open(io.BytesIO(data)) as image:
            image = ImageOps.exif_transpose(image)
            original_size = image.size
            image = self._resize_image(image, max_long_edge)
            out = io.BytesIO()
            if name.endswith((".jpg", ".jpeg")):
                image.convert("RGB").save(out, format="JPEG", quality=jpeg_quality, optimize=True)
            else:
                image.save(out, format="PNG", optimize=True)
            compressed = out.getvalue()
            if len(compressed) >= len(data):
                return data, "kept original; compressed version was not smaller"
            return compressed, f"resized {original_size} -> {image.size}"

    @staticmethod
    def _resize_image(image: Image.Image, max_long_edge: int) -> Image.Image:
        width, height = image.size
        long_edge = max(width, height)
        if long_edge <= max_long_edge:
            return image.copy()
        ratio = max_long_edge / long_edge
        new_size = (max(1, int(width * ratio)), max(1, int(height * ratio)))
        return image.resize(new_size, Image.Resampling.LANCZOS)

    def _compress_scanned_pdf(
        self, blob: bytes, max_long_edge: int, jpeg_quality: int
    ) -> tuple[bytes | None, list[str], str | None]:
        if self._pdf_has_text(blob):
            return None, [], "text PDF kept original to preserve selectable text layer"
        try:
            import pypdfium2 as pdfium
        except Exception as exc:
            return None, [], f"scanned PDF compression skipped; pypdfium2 unavailable: {exc}"

        pdf = pdfium.PdfDocument(io.BytesIO(blob))
        pages: list[Image.Image] = []
        try:
            for page in pdf:
                bitmap = page.render(scale=2).to_pil()
                bitmap = self._resize_image(bitmap.convert("RGB"), max_long_edge)
                pages.append(bitmap)
            if not pages:
                return None, [], "PDF has no pages"
            out = io.BytesIO()
            first, rest = pages[0], pages[1:]
            first.save(out, format="PDF", save_all=True, append_images=rest, quality=jpeg_quality, resolution=144)
            compressed = out.getvalue()
            if len(compressed) >= len(blob):
                return None, [], "scanned PDF compression skipped; compressed version was not smaller"
            return compressed, [f"scanned PDF rasterized to {len(pages)} compressed page images"], None
        finally:
            for page_image in pages:
                page_image.close()
            pdf.close()

    @staticmethod
    def _pdf_has_text(blob: bytes) -> bool:
        try:
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(blob))
            sample_pages = reader.pages[: min(5, len(reader.pages))]
            text = "".join((page.extract_text() or "") for page in sample_pages)
            return len(text.strip()) >= 80
        except Exception:
            return False


    def _parse_fallback_document(self, file: Any, extension: str) -> tuple[str, list[dict[str, Any]]]:
        filename = getattr(file, "filename", "document") or "document"
        blob = getattr(file, "blob", b"") or b""
        if extension in {".xlsx", ".xls"}:
            markdown = self._spreadsheet_to_markdown(blob, extension)
        elif extension == ".csv":
            markdown = self._csv_to_markdown(blob)
        elif extension in {".html", ".htm"}:
            markdown = self._html_to_markdown(blob)
        elif extension in {".json", ".yaml", ".yml"}:
            markdown = "```" + extension.lstrip(".") + "\n" + self._decode_text(blob) + "\n```"
        else:
            markdown = self._decode_text(blob)
        return markdown, [{"filename": filename, "parser": "mmb_fallback_markdown", "extension": extension.lstrip(".")}]

    def _spreadsheet_to_markdown(self, blob: bytes, extension: str) -> str:
        if extension == ".xls":
            import xlrd

            workbook = xlrd.open_workbook(file_contents=blob)
            parts: list[str] = []
            for sheet in workbook.sheets():
                parts.append(f"## Sheet: {sheet.name}")
                rows = [[str(sheet.cell_value(r, c)).strip() for c in range(sheet.ncols)] for r in range(sheet.nrows)]
                parts.append(self._rows_to_markdown(rows))
            return "\n\n".join(parts)

        import openpyxl

        workbook = openpyxl.load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
        parts = []
        for sheet in workbook.worksheets:
            parts.append(f"## Sheet: {sheet.title}")
            rows = [["" if cell is None else str(cell).strip() for cell in row] for row in sheet.iter_rows(values_only=True)]
            parts.append(self._rows_to_markdown(rows))
        workbook.close()
        return "\n\n".join(parts)

    def _csv_to_markdown(self, blob: bytes) -> str:
        text = self._decode_text(blob)
        rows = list(csv.reader(io.StringIO(text)))
        return self._rows_to_markdown(rows)

    @staticmethod
    def _html_to_markdown(blob: bytes) -> str:
        from bs4 import BeautifulSoup

        soup = BeautifulSoup(MmbMaterialPreprocessorTool._decode_text(blob), "html.parser")
        title = soup.title.get_text(strip=True) if soup.title else "HTML Document"
        body = soup.get_text("\n", strip=True)
        return f"# {title}\n\n{body}"

    @staticmethod
    def _rows_to_markdown(rows: list[list[str]]) -> str:
        rows = [row for row in rows if any(cell for cell in row)]
        if not rows:
            return ""
        width = max(len(row) for row in rows)
        normalized = [row + [""] * (width - len(row)) for row in rows]
        header = normalized[0]
        body = normalized[1:]
        lines = ["| " + " | ".join(header) + " |", "| " + " | ".join(["---"] * width) + " |"]
        lines.extend("| " + " | ".join(row) + " |" for row in body[:1000])
        if len(body) > 1000:
            lines.append(f"\n_Only first 1000 rows are indexed from this sheet._")
        return "\n".join(lines)

    @staticmethod
    def _decode_text(blob: bytes) -> str:
        for encoding in ("utf-8-sig", "utf-8", "gb18030", "latin-1"):
            try:
                return blob.decode(encoding)
            except UnicodeDecodeError:
                continue
        return blob.decode("utf-8", errors="replace")

    def _parse_local(
        self, credentials: Credentials, normalized: NormalizedFile, tool_parameters: dict[str, Any]
    ) -> tuple[str, list[dict[str, Any]]]:
        body = {
            "parse_method": tool_parameters.get("parse_method", "auto"),
            "return_md": True,
            "return_model_output": False,
            "return_content_list": True,
            "lang_list": [tool_parameters.get("language") or "ch"],
            "return_images": True,
            "backend": "pipeline",
            "formula_enable": True,
            "table_enable": True,
            "return_middle_json": False,
        }
        file_data = [("files", (normalized.filename, normalized.blob))]
        url = self._api_url(credentials.base_url, "file_parse")
        response = requests.post(url, headers=self._headers(credentials), data=body, files=file_data, timeout=600)
        if self._is_v1_response(response):
            return self._parse_local_v1(credentials, normalized, tool_parameters)
        if response.status_code != 200:
            raise RuntimeError(f"MinerU parse failed: HTTP {response.status_code}, {response.text[:500]}")

        md_parts: list[str] = []
        content_list: list[dict[str, Any]] = []
        results = response.json().get("results", {})
        for file_name, result in results.items():
            images = self._upload_images(result.get("images") or {})
            if result.get("content_list"):
                try:
                    parsed = json.loads(result["content_list"])
                    if isinstance(parsed, list):
                        content_list.extend(parsed)
                    else:
                        content_list.append(parsed)
                except json.JSONDecodeError:
                    content_list.append({"filename": file_name, "raw_content_list": result.get("content_list")})
            md_content = result.get("md_content") or ""
            md_parts.append(self._replace_md_img_path(md_content, images))
        return "\n\n".join(part for part in md_parts if part), content_list

    def _parse_local_v1(
        self, credentials: Credentials, normalized: NormalizedFile, tool_parameters: dict[str, Any]
    ) -> tuple[str, list[dict[str, Any]]]:
        params = {
            "parse_method": tool_parameters.get("parse_method", "auto"),
            "return_layout": False,
            "return_info": False,
            "return_content_list": True,
            "return_images": True,
        }
        file_data = {"file": (normalized.filename, normalized.blob)}
        url = self._api_url(credentials.base_url, "file_parse")
        response = requests.post(url, headers=self._headers(credentials), params=params, files=file_data, timeout=600)
        if response.status_code != 200:
            raise RuntimeError(f"MinerU v1 parse failed: HTTP {response.status_code}, {response.text[:500]}")
        payload = response.json()
        images = self._upload_images(payload.get("images") or {})
        md_content = self._replace_md_img_path(payload.get("md_content") or "", images)
        content_list = payload.get("content_list") or []
        if not isinstance(content_list, list):
            content_list = [content_list]
        return md_content, content_list

    @staticmethod
    def _is_v1_response(response: requests.Response) -> bool:
        if response.status_code != 422:
            return False
        try:
            detail = response.json().get("detail")
        except Exception:
            return False
        if not isinstance(detail, list):
            return False
        return any(item.get("loc", [])[:2] == ["body", "file"] for item in detail if isinstance(item, dict))

    def _upload_images(self, file_obj: dict[str, str]) -> list[Any]:
        images = []
        import base64

        for file_name, encoded in file_obj.items():
            try:
                base64_data = encoded.split(",", 1)[1] if "," in encoded else encoded
                image_bytes = base64.b64decode(base64_data)
                mime_type = mimetypes.guess_type(file_name)[0] or "image/jpeg"
                images.append(self.session.file.upload(file_name, image_bytes, mime_type))
            except Exception as exc:
                logger.warning("Failed to upload extracted image %s: %s", file_name, exc)
        return images

    @staticmethod
    def _replace_md_img_path(md_content: str, images: list[Any]) -> str:
        for image in images:
            preview_url = getattr(image, "preview_url", None)
            name = getattr(image, "name", None)
            if preview_url and name:
                md_content = md_content.replace(f"images/{name}", preview_url)
        return md_content

    @staticmethod
    def _build_file_metadata(file: Any, normalized: NormalizedFile, extension: str) -> dict[str, Any]:
        filename = getattr(file, "filename", "document") or "document"
        return {
            "source_file_name": filename,
            "normalized_file_name": normalized.filename,
            "file_extension": extension.lstrip("."),
            "parser": "mmb_material_preprocessor",
            "original_size_bytes": normalized.report["original_size_bytes"],
            "normalized_size_bytes": normalized.report["normalized_size_bytes"],
            "compression_applied": normalized.report["compressed"],
            "image_strategy": "mineru_ocr_with_vision_fallback_ocr_only",
        }

    @staticmethod
    def _with_metadata_header(md_content: str, metadata: dict[str, Any], report: dict[str, Any]) -> str:
        metadata_lines = ["# 文件解析元数据"]
        for key, value in metadata.items():
            metadata_lines.append(f"- {key}: {value}")
        metadata_lines.append("\n# 压缩与解析报告")
        metadata_lines.append(f"```json\n{json.dumps(report, ensure_ascii=False, indent=2)}\n```")
        metadata_lines.append("\n# 解析正文")
        return "\n".join(metadata_lines) + "\n\n" + (md_content or "")
