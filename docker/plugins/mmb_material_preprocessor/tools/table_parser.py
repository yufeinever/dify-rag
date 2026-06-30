from __future__ import annotations

import csv
import io
import json
import os
import re
from dataclasses import dataclass
from html.parser import HTMLParser
from typing import Any, Sequence


SUPPORTED_TABLE_EXTENSIONS = {".csv", ".xls", ".xlsx"}
_KS_SOURCE_RE = re.compile(r"#(?P<sheet>.+)!(?P<start_col>[A-Z]+)\d+:(?P<end_col>[A-Z]+)\d+$")


@dataclass
class TableParseResult:
    text: str
    table_metadata: dict[str, Any]
    parse_report: dict[str, Any]
    content_list: list[dict[str, Any]]


def parse_table_file(filename: str, blob: bytes) -> TableParseResult:
    extension = os.path.splitext(filename)[1].lower()
    if extension not in SUPPORTED_TABLE_EXTENSIONS:
        raise ValueError(f"Unsupported table extension {extension}. Supported extensions: {sorted(SUPPORTED_TABLE_EXTENSIONS)}")

    if extension == ".csv":
        documents, content_list, sheet_stats = _parse_csv(filename, blob)
        parser = "csv"
    elif extension == ".xlsx":
        documents, content_list, sheet_stats, parser = _parse_xlsx(filename, blob)
    elif _looks_like_html(blob):
        documents, content_list, sheet_stats = _parse_html_xls(filename, blob)
        parser = "html-table"
    else:
        documents, content_list, sheet_stats = _parse_xls(filename, blob)
        parser = "xlrd"

    metadata = {
        "source_file_name": filename,
        "file_extension": extension.lstrip("."),
        "parser": parser,
        "sheet_count": len(sheet_stats),
        "sheets": sheet_stats,
    }
    report = {
        "parser": parser,
        "row_documents": len([item for item in content_list if item.get("type") == "row"]),
        "table_documents": len([item for item in content_list if item.get("type") == "table"]),
        "warnings": [],
    }
    text = _with_metadata_header("\n\n".join(documents), metadata, report)
    return TableParseResult(text=text, table_metadata=metadata, parse_report=report, content_list=content_list)


def _parse_xlsx(filename: str, blob: bytes) -> tuple[list[str], list[dict[str, Any]], list[dict[str, Any]], str]:
    source_ranges: dict[str, tuple[str, str, str]] = {}
    table_summaries: list[str] = []
    content_list: list[dict[str, Any]] = []
    parser = "openpyxl"

    try:
        from ks_xlsx_parser import parse_workbook

        result = parse_workbook(content=blob, filename=filename or "workbook.xlsx")
        chunks = list(getattr(result, "chunks", []) or [])
        if chunks:
            parser = "ks-xlsx-parser"
            source_ranges = _ks_source_ranges(chunks)
            for index, chunk in enumerate(chunks, start=1):
                summary_text, summary_item = _format_ks_table_summary(filename, index, chunk)
                if summary_text:
                    table_summaries.append(summary_text)
                if summary_item:
                    content_list.append(summary_item)
    except Exception as exc:
        table_summaries.append(f"Parser warning: ks-xlsx-parser unavailable, fallback to openpyxl ({exc})")

    row_documents, row_items, sheet_stats = _xlsx_rows_with_openpyxl(filename, blob, source_ranges)
    return [*table_summaries, *row_documents], [*content_list, *row_items], sheet_stats, parser


def _xlsx_rows_with_openpyxl(
    filename: str, blob: bytes, source_ranges: dict[str, tuple[str, str, str]]
) -> tuple[list[str], list[dict[str, Any]], list[dict[str, Any]]]:
    import openpyxl

    workbook = openpyxl.load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    documents: list[str] = []
    content_list: list[dict[str, Any]] = []
    sheet_stats: list[dict[str, Any]] = []
    try:
        for sheet in workbook.worksheets:
            rows = [[_clean_value(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
            sheet_docs, sheet_items, stats = _rows_to_documents(
                filename=filename,
                sheet_name=sheet.title,
                rows=rows,
                source_ranges=source_ranges,
            )
            documents.extend(sheet_docs)
            content_list.extend(sheet_items)
            sheet_stats.append(stats)
    finally:
        workbook.close()
    return documents, content_list, sheet_stats


def _parse_xls(filename: str, blob: bytes) -> tuple[list[str], list[dict[str, Any]], list[dict[str, Any]]]:
    import xlrd

    workbook = xlrd.open_workbook(file_contents=blob)
    documents: list[str] = []
    content_list: list[dict[str, Any]] = []
    sheet_stats: list[dict[str, Any]] = []
    for sheet in workbook.sheets():
        rows = [[_clean_value(sheet.cell_value(r, c)) for c in range(sheet.ncols)] for r in range(sheet.nrows)]
        sheet_docs, sheet_items, stats = _rows_to_documents(filename=filename, sheet_name=sheet.name, rows=rows)
        documents.extend(sheet_docs)
        content_list.extend(sheet_items)
        sheet_stats.append(stats)
    return documents, content_list, sheet_stats


def _parse_csv(filename: str, blob: bytes) -> tuple[list[str], list[dict[str, Any]], list[dict[str, Any]]]:
    text = _decode_text(blob)
    rows = list(csv.reader(io.StringIO(text)))
    documents, content_list, stats = _rows_to_documents(filename=filename, sheet_name="CSV", rows=rows)
    return documents, content_list, [stats]


def _parse_html_xls(filename: str, blob: bytes) -> tuple[list[str], list[dict[str, Any]], list[dict[str, Any]]]:
    parser = _HTMLTableParser()
    parser.feed(_decode_text(blob))
    documents: list[str] = []
    content_list: list[dict[str, Any]] = []
    sheet_stats: list[dict[str, Any]] = []
    for index, table in enumerate(parser.tables, start=1):
        sheet_name = f"HTML Table {index}"
        sheet_docs, sheet_items, stats = _rows_to_documents(filename=filename, sheet_name=sheet_name, rows=table)
        documents.extend(sheet_docs)
        content_list.extend(sheet_items)
        sheet_stats.append(stats)
    return documents, content_list, sheet_stats


def _rows_to_documents(
    *,
    filename: str,
    sheet_name: str,
    rows: Sequence[Sequence[Any]],
    source_ranges: dict[str, tuple[str, str, str]] | None = None,
) -> tuple[list[str], list[dict[str, Any]], dict[str, Any]]:
    indexed_rows = [(idx, [_clean_value(cell) for cell in row]) for idx, row in enumerate(rows, start=1)]
    indexed_rows = [(idx, row) for idx, row in indexed_rows if any(row)]
    if not indexed_rows:
        return [], [], {"sheet_name": sheet_name, "data_rows": 0, "source": ""}

    header_number, header_row = indexed_rows[0]
    headers = [cell or f"Column {idx + 1}" for idx, cell in enumerate(header_row)]
    max_col = max(len(row) for _, row in indexed_rows)
    source = _table_source_uri(filename, sheet_name, header_number, indexed_rows[-1][0], max_col, source_ranges)

    summary = [
        f"File: {filename}",
        f"Sheet: {sheet_name}",
        f"sheet={sheet_name}",
        f"Table rows: {max(0, len(indexed_rows) - 1)}",
    ]
    if source:
        summary.append(f"Source: {source}")

    documents = ["\n".join(summary)]
    content_list = [
        {
            "type": "table",
            "filename": filename,
            "sheet": sheet_name,
            "source": source,
            "data_rows": max(0, len(indexed_rows) - 1),
            "headers": headers,
        }
    ]

    for row_number, row in indexed_rows[1:]:
        fields: list[tuple[str, str]] = []
        for index, header in enumerate(headers):
            value = row[index] if index < len(row) else ""
            if value:
                fields.append((header, value))
        if not fields:
            continue
        source_uri = _row_source_uri(filename, sheet_name, row_number, max(len(headers), len(row)), source_ranges)
        documents.append(_format_row(filename, sheet_name, row_number, fields, source_uri))
        content_list.append(
            {
                "type": "row",
                "filename": filename,
                "sheet": sheet_name,
                "row": row_number,
                "source": source_uri,
                "fields": dict(fields),
            }
        )

    stats = {"sheet_name": sheet_name, "data_rows": len(content_list) - 1, "source": source}
    return documents, content_list, stats


def _format_ks_table_summary(filename: str, table_index: int, chunk: Any) -> tuple[str, dict[str, Any]]:
    source_uri = _clean_value(getattr(chunk, "source_uri", ""))
    block_type = _clean_value(getattr(chunk, "block_type", ""))
    token_count = getattr(chunk, "token_count", None)
    render_text = _clean_value(getattr(chunk, "render_text", ""))
    sheet_name = _sheet_from_source(source_uri)
    summary = render_text.split("|")[0].strip() if render_text else ""

    lines = [f"File: {filename}", f"Table: {table_index}", "Parser: ks-xlsx-parser"]
    if sheet_name:
        lines.extend([f"Sheet: {sheet_name}", f"sheet={sheet_name}"])
    if source_uri:
        lines.append(f"Source: {source_uri}")
    if block_type:
        lines.append(f"Block: {block_type}")
    if token_count is not None:
        lines.append(f"Tokens: {token_count}")
    if summary:
        lines.append(f"Summary: {summary}")

    return "\n".join(lines), {
        "type": "table",
        "filename": filename,
        "sheet": sheet_name,
        "source": source_uri,
        "parser": "ks-xlsx-parser",
        "block_type": block_type,
        "token_count": token_count,
        "summary": summary,
    }


def _format_row(
    filename: str, sheet_name: str, row_number: int, fields: Sequence[tuple[str, str]], source_uri: str = ""
) -> str:
    lines = [f"File: {filename}", f"Sheet: {sheet_name}", f"sheet={sheet_name}", f"Row: {row_number}"]
    if source_uri:
        lines.append(f"Source: {source_uri}")
    lines.extend(f"{name}: {value}" for name, value in fields)
    return "\n".join(lines)


def _ks_source_ranges(chunks: Sequence[Any]) -> dict[str, tuple[str, str, str]]:
    ranges: dict[str, tuple[str, str, str]] = {}
    for chunk in chunks:
        source_uri = _clean_value(getattr(chunk, "source_uri", ""))
        match = _KS_SOURCE_RE.search(source_uri)
        if not match:
            continue
        ranges[match.group("sheet")] = (source_uri.split("#", 1)[0], match.group("start_col"), match.group("end_col"))
    return ranges


def _table_source_uri(
    filename: str,
    sheet_name: str,
    first_row: int,
    last_row: int,
    max_col: int,
    source_ranges: dict[str, tuple[str, str, str]] | None,
) -> str:
    if source_ranges and sheet_name in source_ranges:
        prefix, start_col, end_col = source_ranges[sheet_name]
        return f"{prefix or filename}#{sheet_name}!{start_col}{first_row}:{end_col}{last_row}"
    return f"{filename}#{sheet_name}!A{first_row}:{_column_name(max_col)}{last_row}"


def _row_source_uri(
    filename: str,
    sheet_name: str,
    row_number: int,
    max_col: int,
    source_ranges: dict[str, tuple[str, str, str]] | None,
) -> str:
    if source_ranges and sheet_name in source_ranges:
        prefix, start_col, end_col = source_ranges[sheet_name]
        return f"{prefix or filename}#{sheet_name}!{start_col}{row_number}:{end_col}{row_number}"
    return f"{filename}#{sheet_name}!A{row_number}:{_column_name(max_col)}{row_number}"


def _column_name(index: int) -> str:
    result = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        result = chr(65 + remainder) + result
    return result or "A"


def _sheet_from_source(source_uri: str) -> str:
    if "#" not in source_uri or "!" not in source_uri:
        return ""
    return source_uri.split("#", 1)[1].split("!", 1)[0]


def _with_metadata_header(content: str, metadata: dict[str, Any], report: dict[str, Any]) -> str:
    lines = ["# 表格解析元数据"]
    for key, value in metadata.items():
        if key == "sheets":
            lines.append(f"- sheets: {json.dumps(value, ensure_ascii=False)}")
        else:
            lines.append(f"- {key}: {value}")
    lines.append("\n# 表格解析报告")
    lines.append(f"```json\n{json.dumps(report, ensure_ascii=False, indent=2)}\n```")
    lines.append("\n# 表格解析正文")
    return "\n".join(lines) + "\n\n" + (content or "")


def _looks_like_html(blob: bytes) -> bool:
    prefix = _decode_text(blob[:4096]).lstrip().lower()
    return prefix.startswith(("<!doctype html", "<html", "<table"))


def _decode_text(blob: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "latin-1"):
        try:
            return blob.decode(encoding)
        except UnicodeDecodeError:
            continue
    return blob.decode("utf-8", errors="replace")


def _clean_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return " ".join(str(value).split())


class _HTMLTableParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.tables: list[list[list[str]]] = []
        self._table_stack = 0
        self._current_table: list[list[str]] | None = None
        self._current_row: list[str] | None = None
        self._current_cell: list[str] | None = None

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag == "table":
            self._table_stack += 1
            if self._table_stack == 1:
                self._current_table = []
        elif tag == "tr" and self._table_stack == 1:
            self._current_row = []
        elif tag in {"td", "th"} and self._table_stack == 1 and self._current_row is not None:
            self._current_cell = []

    def handle_data(self, data: str) -> None:
        if self._current_cell is not None:
            self._current_cell.append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag in {"td", "th"} and self._current_cell is not None and self._current_row is not None:
            self._current_row.append(_clean_value("".join(self._current_cell)))
            self._current_cell = None
        elif tag == "tr" and self._current_row is not None and self._current_table is not None:
            self._current_table.append(self._current_row)
            self._current_row = None
        elif tag == "table" and self._table_stack:
            if self._table_stack == 1 and self._current_table is not None:
                self.tables.append(self._current_table)
                self._current_table = None
            self._table_stack -= 1
