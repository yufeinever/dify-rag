from __future__ import annotations

import io
import json
import os
import re
from dataclasses import dataclass
from typing import Any

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches as PptInches
from pptx.util import Pt as PptPt


DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
MAX_MARKDOWN_CHARS = 80000
MAX_PPT_CHARS = 60000
MAX_SLIDES = 60


@dataclass
class Artifact:
    filename: str
    blob: bytes
    mime_type: str
    summary: dict[str, Any]


def safe_filename(name: str | None, default_stem: str, extension: str) -> str:
    raw = (name or default_stem).strip() or default_stem
    raw = raw.replace("\\", "_").replace("/", "_")
    raw = re.sub(r"[\x00-\x1f<>:\"|?*]+", "_", raw)
    raw = re.sub(r"\s+", " ", raw).strip(" .") or default_stem
    if not raw.lower().endswith(extension):
        raw += extension
    return raw[:120]


def _set_east_asia_font(style: Any, font_name: str = "Microsoft YaHei") -> None:
    if style._element.rPr is None:
        style._element.get_or_add_rPr()
    style._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def _set_cell_shading(cell: Any, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_table_geometry(table: Any, widths: list[int]) -> None:
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = width
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")


def _configure_doc_styles(doc: Document, style_preset: str) -> None:
    styles = doc.styles
    base_font = "Arial" if style_preset == "google_docs" else "Calibri"
    heading_color = "000000" if style_preset == "google_docs" else "2E74B5"
    normal = styles["Normal"]
    normal.font.name = base_font
    normal.font.size = Pt(11)
    _set_east_asia_font(normal)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15 if style_preset == "google_docs" else 1.1
    for name, size, color in [
        ("Heading 1", 16, heading_color),
        ("Heading 2", 13, heading_color),
        ("Heading 3", 12, "1F4D78" if style_preset != "google_docs" else "434343"),
    ]:
        style = styles[name]
        style.font.name = base_font
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        _set_east_asia_font(style)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)


def _split_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def _is_table_separator(line: str) -> bool:
    cells = _split_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", c or "") for c in cells)


def _add_doc_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    rows = [row + [""] * (column_count - len(row)) for row in rows]
    table = doc.add_table(rows=1, cols=column_count)
    table.style = "Table Grid"
    widths = [max(1200, int(9360 / column_count))] * column_count
    widths[-1] += 9360 - sum(widths)
    _set_table_geometry(table, widths)
    for idx, value in enumerate(rows[0]):
        table.rows[0].cells[idx].text = value
        _set_cell_shading(table.rows[0].cells[idx], "F2F4F7")
        for paragraph in table.rows[0].cells[idx].paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.bold = True
    for row in rows[1:]:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
    doc.add_paragraph()


def build_docx_artifact(
    *,
    title: str,
    markdown_content: str,
    filename: str | None = None,
    style_preset: str = "business_brief",
) -> Artifact:
    content = (markdown_content or "").strip()
    if not content:
        raise ValueError("markdown_content is required")
    if len(content) > MAX_MARKDOWN_CHARS:
        raise ValueError(f"markdown_content is too long; max {MAX_MARKDOWN_CHARS} characters")

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
    _configure_doc_styles(doc, style_preset)

    title_para = doc.add_paragraph()
    title_para.paragraph_format.space_after = Pt(12)
    title_run = title_para.add_run((title or "生成文档").strip())
    title_run.bold = True
    title_run.font.size = Pt(22 if style_preset != "google_docs" else 20)
    title_run.font.color.rgb = RGBColor.from_string("0B2545" if style_preset != "google_docs" else "000000")

    lines = content.splitlines()
    i = 0
    table_count = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("|") and "|" in stripped[1:]:
            table_rows: list[list[str]] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = _split_table_row(lines[i])
                if not _is_table_separator(lines[i]):
                    table_rows.append(cells)
                i += 1
            _add_doc_table(doc, table_rows)
            table_count += 1
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading:
            doc.add_heading(heading.group(2).strip(), level=len(heading.group(1)))
        elif re.match(r"^[-*+]\s+", stripped):
            doc.add_paragraph(re.sub(r"^[-*+]\s+", "", stripped), style="List Bullet")
        elif re.match(r"^\d+[.)]\s+", stripped):
            doc.add_paragraph(re.sub(r"^\d+[.)]\s+", "", stripped), style="List Number")
        else:
            paragraph_lines = [stripped]
            while i + 1 < len(lines):
                nxt = lines[i + 1].strip()
                if not nxt or nxt.startswith("#") or re.match(r"^([-*+]|\d+[.)])\s+", nxt) or nxt.startswith("|"):
                    break
                paragraph_lines.append(nxt)
                i += 1
            doc.add_paragraph(" ".join(paragraph_lines))
        i += 1

    out = io.BytesIO()
    doc.save(out)
    blob = out.getvalue()
    return Artifact(
        filename=safe_filename(filename, title or "document", ".docx"),
        blob=blob,
        mime_type=DOCX_MIME,
        summary={"format": "docx", "size_bytes": len(blob), "tables": table_count, "characters": len(content)},
    )


def _slides_from_json(slides_json: str) -> list[dict[str, Any]]:
    data = json.loads(slides_json)
    if not isinstance(data, list):
        raise ValueError("slides_json must be a JSON array")
    slides: list[dict[str, Any]] = []
    for item in data:
        if not isinstance(item, dict):
            continue
        title = str(item.get("title") or "").strip()
        bullets = item.get("bullets") or []
        if isinstance(bullets, str):
            bullets = [line.strip() for line in bullets.splitlines() if line.strip()]
        slides.append({"title": title or "Untitled", "bullets": [str(b).strip() for b in bullets if str(b).strip()]})
    return slides


def _slides_from_markdown(title: str, markdown_outline: str) -> list[dict[str, Any]]:
    slides: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None
    for raw in markdown_outline.splitlines():
        line = raw.strip()
        if not line:
            continue
        if line == "---":
            current = None
            continue
        heading = re.match(r"^#{1,3}\s+(.+)$", line)
        if heading:
            current = {"title": heading.group(1).strip(), "bullets": []}
            slides.append(current)
            continue
        bullet = re.match(r"^[-*+]\s+(.+)$", line)
        if bullet:
            if current is None:
                current = {"title": title or "方案要点", "bullets": []}
                slides.append(current)
            current["bullets"].append(bullet.group(1).strip())
        else:
            if current is None:
                current = {"title": line[:36], "bullets": []}
                slides.append(current)
            else:
                current["bullets"].append(line)
    return slides or [{"title": title or "方案", "bullets": [markdown_outline.strip()[:500]]}]


def _split_dense_slides(slides: list[dict[str, Any]]) -> list[dict[str, Any]]:
    expanded: list[dict[str, Any]] = []
    for slide in slides:
        bullets = slide.get("bullets") or []
        if len(bullets) <= 7:
            expanded.append(slide)
            continue
        for idx in range(0, len(bullets), 7):
            suffix = "" if idx == 0 else f"（续 {idx // 7 + 1}）"
            expanded.append({"title": f"{slide.get('title', 'Untitled')}{suffix}", "bullets": bullets[idx : idx + 7]})
    return expanded[:MAX_SLIDES]


def build_pptx_artifact(
    *,
    title: str,
    markdown_outline: str = "",
    slides_json: str = "",
    filename: str | None = None,
    theme: str = "mmb_business",
) -> Artifact:
    source = (slides_json or markdown_outline or "").strip()
    if not source:
        raise ValueError("slides_json or markdown_outline is required")
    if len(source) > MAX_PPT_CHARS:
        raise ValueError(f"PPT content is too long; max {MAX_PPT_CHARS} characters")

    slides = _slides_from_json(slides_json) if slides_json.strip() else _slides_from_markdown(title, markdown_outline)
    slides = _split_dense_slides(slides)

    prs = Presentation()
    prs.slide_width = PptInches(13.333)
    prs.slide_height = PptInches(7.5)
    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = title or "MMB 方案"
    title_slide.placeholders[1].text = "Generated by MMB Office Artifact Tools"

    for slide_data in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = slide_data.get("title") or "Untitled"
        body = slide.placeholders[1].text_frame
        body.clear()
        bullets = slide_data.get("bullets") or [" "]
        for idx, bullet in enumerate(bullets[:7]):
            paragraph = body.paragraphs[0] if idx == 0 else body.add_paragraph()
            paragraph.text = bullet[:220]
            paragraph.level = 0
            paragraph.font.size = PptPt(22 if len(bullets) <= 5 else 19)
        for shape in slide.shapes:
            if hasattr(shape, "text_frame"):
                for paragraph in shape.text_frame.paragraphs:
                    paragraph.alignment = PP_ALIGN.LEFT

    out = io.BytesIO()
    prs.save(out)
    blob = out.getvalue()
    return Artifact(
        filename=safe_filename(filename, title or "deck", ".pptx"),
        blob=blob,
        mime_type=PPTX_MIME,
        summary={"format": "pptx", "size_bytes": len(blob), "slides": len(prs.slides), "theme": theme},
    )
