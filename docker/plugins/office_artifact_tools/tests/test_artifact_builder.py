import io
import unittest

from docx import Document
from pptx import Presentation
from openpyxl import load_workbook

from tools.artifact_builder import build_docx_artifact, build_pptx_artifact, build_xlsx_artifact, safe_filename


class ArtifactBuilderTests(unittest.TestCase):
    def test_safe_filename_adds_extension_and_removes_bad_chars(self):
        self.assertEqual(safe_filename('bad/name:*?', 'doc', '.docx'), 'bad_name_.docx')

    def test_build_docx_from_markdown(self):
        artifact = build_docx_artifact(
            title='MMB SaaS 方案',
            filename='mmb方案.docx',
            markdown_content='''# 结论\nMMB 应定位为设备运营系统。\n\n- 设备可靠性\n- 支付对账\n\n| 模块 | 价值 |\n| --- | --- |\n| 设备 | 在线监控 |\n''',
        )
        self.assertTrue(artifact.filename.endswith('.docx'))
        self.assertGreater(len(artifact.blob), 10000)
        self.assertEqual(artifact.summary['tables'], 1)
        doc = Document(io.BytesIO(artifact.blob))
        self.assertGreaterEqual(len(doc.paragraphs), 4)
        self.assertEqual(len(doc.tables), 1)

    def test_build_pptx_from_markdown_outline(self):
        artifact = build_pptx_artifact(
            title='MMB SaaS 方案',
            markdown_outline='''# 定位\n- 鲜啤设备运营系统\n- 支付和运维闭环\n\n# MVP\n- 设备可靠性\n- 订单对账\n''',
        )
        self.assertTrue(artifact.filename.endswith('.pptx'))
        self.assertGreater(len(artifact.blob), 10000)
        deck = Presentation(io.BytesIO(artifact.blob))
        self.assertEqual(len(deck.slides), 3)

    def test_build_xlsx_from_sheets_json(self):
        artifact = build_xlsx_artifact(
            title='MMB 端午预算表',
            filename='端午预算.xlsx',
            sheets_json='''[
                {
                    "name": "预算明细",
                    "columns": ["项目", "预算", "负责人"],
                    "rows": [["端午限定物料", 8000, "市场部"], ["线下快闪", 20000, "运营部"]]
                },
                {
                    "name": "排期",
                    "columns": ["日期", "动作"],
                    "rows": [{"日期": "6月1日", "动作": "预热"}]
                }
            ]''',
        )
        self.assertTrue(artifact.filename.endswith('.xlsx'))
        self.assertGreater(len(artifact.blob), 4000)
        self.assertEqual(artifact.summary['sheets'], 2)
        workbook = load_workbook(io.BytesIO(artifact.blob))
        self.assertEqual(workbook.sheetnames, ['预算明细', '排期'])
        self.assertEqual(workbook['预算明细']['A2'].value, '端午限定物料')
        self.assertEqual(workbook['预算明细']['B2'].value, 8000)

    def test_build_xlsx_from_markdown_table(self):
        artifact = build_xlsx_artifact(
            title='简单预算',
            table_markdown='| 项目 | 预算 |\n| --- | --- |\n| 物料 | 8000 |',
        )
        self.assertEqual(artifact.summary['sheets'], 1)
        workbook = load_workbook(io.BytesIO(artifact.blob))
        self.assertEqual(workbook.active['A2'].value, '物料')

    def test_build_xlsx_from_fenced_json(self):
        artifact = build_xlsx_artifact(
            title='端午预算',
            sheets_json='''```json
[
  {"name": "预算明细", "headers": ["项目", "预算"], "data": [["物料", 8000]]}
]
```''',
        )
        workbook = load_workbook(io.BytesIO(artifact.blob))
        self.assertEqual(workbook.sheetnames, ['预算明细'])
        self.assertEqual(workbook['预算明细']['B2'].value, 8000)

    def test_build_xlsx_from_double_encoded_json(self):
        artifact = build_xlsx_artifact(
            title='端午预算',
            sheets_json='''"[{\\"name\\": \\"排期\\", \\"columns\\": [\\"日期\\", \\"动作\\"], \\"rows\\": [[\\"6月1日\\", \\"预热\\"]]}]"''',
        )
        workbook = load_workbook(io.BytesIO(artifact.blob))
        self.assertEqual(workbook.sheetnames, ['排期'])
        self.assertEqual(workbook['排期']['B2'].value, '预热')

    def test_build_xlsx_from_multi_markdown_tables(self):
        artifact = build_xlsx_artifact(
            title='端午营销',
            content='''# 预算明细

| 项目 | 预算 |
| --- | --- |
| 物料 | 8000 |

# 活动排期

| 日期 | 动作 |
| --- | --- |
| 6月1日 | 预热 |
''',
        )
        workbook = load_workbook(io.BytesIO(artifact.blob))
        self.assertEqual(workbook.sheetnames, ['预算明细', '活动排期'])
        self.assertEqual(workbook['预算明细']['A2'].value, '物料')
        self.assertEqual(workbook['活动排期']['B2'].value, '预热')

    def test_empty_docx_content_is_rejected(self):
        with self.assertRaises(ValueError):
            build_docx_artifact(title='Empty', markdown_content='')

    def test_empty_ppt_content_is_rejected(self):
        with self.assertRaises(ValueError):
            build_pptx_artifact(title='Empty')

    def test_empty_xlsx_content_is_rejected(self):
        with self.assertRaises(ValueError):
            build_xlsx_artifact(title='Empty')


if __name__ == '__main__':
    unittest.main()
