import io
import json
import unittest

from PIL import Image
from pptx import Presentation

from tools.visual_ppt_builder import (
    build_prompt_for_slide,
    build_visual_ppt_artifact,
    normalize_slide_image,
    slides_from_markdown,
    slides_from_json,
)


def fake_png(width=800, height=600, color=(250, 128, 40)):
    image = Image.new("RGB", (width, height), color=color)
    out = io.BytesIO()
    image.save(out, format="PNG")
    return out.getvalue()


def fake_image_client(prompt, model, quality, size):
    return fake_png()


class VisualPptBuilderTest(unittest.TestCase):
    def test_markdown_outline_to_slides(self):
        slides = slides_from_markdown(
            "端午营销",
            """
# 封面
- MMB 端午节
- 鲜啤消费场景

## 活动策略
- 广场夜饮
- 自助设备
""",
            slide_count=6,
        )
        self.assertEqual(len(slides), 2)
        self.assertEqual(slides[0].title, "封面")
        self.assertIn("鲜啤消费场景", slides[0].bullets)

    def test_slides_json_has_priority_shape(self):
        slides = slides_from_json(
            json.dumps(
                [
                    {"title": "封面", "bullets": ["MMB", "端午"], "image_prompt": "hero visual"},
                    {"title": "方案", "bullets": "设备+场景"},
                ],
                ensure_ascii=False,
            ),
            slide_count=6,
        )
        self.assertEqual(len(slides), 2)
        self.assertEqual(slides[0].image_prompt, "hero visual")
        self.assertEqual(slides[1].bullets, ["设备+场景"])

    def test_normalize_image_to_16_9_png(self):
        normalized = normalize_slide_image(fake_png(800, 600))
        image = Image.open(io.BytesIO(normalized))
        self.assertEqual(image.size, (1536, 864))

    def test_build_prompt_is_visual_slide_prompt(self):
        slide = slides_from_markdown("MMB", "# 商业机会\n- 自助鲜啤\n- 夜间消费", 6)[0]
        prompt = build_prompt_for_slide(slide, deck_title="MMB", style_preset="mmb_modern_pitch")
        self.assertIn("16:9", prompt)
        self.assertIn("商业机会", prompt)
        self.assertIn("avoid dense small text", prompt)

    def test_build_visual_ppt_with_mocked_images(self):
        artifact = build_visual_ppt_artifact(
            title="端午营销视觉PPT",
            outline="# 封面\n- MMB 端午节\n\n# 策略\n- 广场夜饮",
            filename="端午营销视觉PPT.pptx",
            slide_count=6,
            image_client=fake_image_client,
        )
        self.assertTrue(artifact.filename.endswith(".pptx"))
        self.assertEqual(artifact.mime_type, "application/vnd.openxmlformats-officedocument.presentationml.presentation")
        prs = Presentation(io.BytesIO(artifact.blob))
        self.assertEqual(len(prs.slides), 2)
        for slide in prs.slides:
            self.assertGreaterEqual(len(slide.shapes), 1)
        self.assertEqual(artifact.summary["slides"], 2)

    def test_rejects_empty_outline(self):
        with self.assertRaises(ValueError):
            build_visual_ppt_artifact(title="空", outline="", image_client=fake_image_client)

    def test_rejects_bad_json(self):
        with self.assertRaises(ValueError):
            slides_from_json("{bad", slide_count=6)

    def test_caps_slide_count(self):
        outline = "\n".join([f"# 第{i}页\n- 内容" for i in range(20)])
        artifact = build_visual_ppt_artifact(
            title="多页",
            outline=outline,
            slide_count=20,
            image_client=fake_image_client,
        )
        prs = Presentation(io.BytesIO(artifact.blob))
        self.assertEqual(len(prs.slides), 12)


if __name__ == "__main__":
    unittest.main()
