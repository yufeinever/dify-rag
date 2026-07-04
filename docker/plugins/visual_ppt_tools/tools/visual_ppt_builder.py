from __future__ import annotations

import base64
import hashlib
import io
import json
import os
import re
import tempfile
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable
from urllib.parse import urljoin

import requests
from PIL import Image
from pptx import Presentation
from pptx.util import Inches


PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
MAX_OUTLINE_CHARS = 30000
MAX_SLIDES = 12
DEFAULT_SLIDE_COUNT = 6
TARGET_IMAGE_SIZE = (1536, 864)
DEFAULT_IMAGE_CONCURRENCY = 6
REQUESTED_IMAGE_SIZE = "1536x864"
FALLBACK_IMAGE_SIZE = "1536x1024"


@dataclass
class Artifact:
    filename: str
    blob: bytes
    mime_type: str
    summary: dict[str, Any]


@dataclass
class VisualSlide:
    title: str
    bullets: list[str]
    image_prompt: str = ""


@dataclass
class OpenAIImageConfig:
    api_key: str
    api_base: str = "https://api.openai.com/v1"
    model: str = "gpt-image-2"
    quality: str = "medium"
    size: str = REQUESTED_IMAGE_SIZE


ImageClient = Callable[[str, str, str, str], bytes]


def safe_filename(name: str | None, default_stem: str, extension: str) -> str:
    raw = (name or default_stem).strip() or default_stem
    raw = raw.replace("\\", "_").replace("/", "_")
    raw = re.sub(r"[\x00-\x1f<>:\"|?*]+", "_", raw)
    raw = re.sub(r"\s+", " ", raw).strip(" .") or default_stem
    if not raw.lower().endswith(extension):
        raw += extension
    return raw[:120]


def normalize_slide_count(value: Any) -> int:
    try:
        count = int(value or DEFAULT_SLIDE_COUNT)
    except (TypeError, ValueError):
        count = DEFAULT_SLIDE_COUNT
    return max(1, min(MAX_SLIDES, count))


def _clean_text(value: Any, limit: int = 120) -> str:
    text = re.sub(r"\s+", " ", str(value or "")).strip()
    return text[:limit]


def _normalize_bullets(raw: Any) -> list[str]:
    if raw is None:
        return []
    if isinstance(raw, str):
        lines = [line.strip(" -\t") for line in raw.splitlines() if line.strip()]
        return [_clean_text(line, 80) for line in lines if _clean_text(line, 80)][:4]
    if isinstance(raw, list):
        bullets = [_clean_text(item, 80) for item in raw]
        return [item for item in bullets if item][:4]
    return [_clean_text(raw, 80)] if _clean_text(raw, 80) else []


def slides_from_json(slides_json: str, slide_count: Any = DEFAULT_SLIDE_COUNT) -> list[VisualSlide]:
    try:
        payload = json.loads(slides_json)
    except json.JSONDecodeError as exc:
        raise ValueError(f"slides_json is invalid JSON: {exc.msg}") from exc
    if not isinstance(payload, list):
        raise ValueError("slides_json must be an array")

    limit = normalize_slide_count(slide_count)
    slides: list[VisualSlide] = []
    for index, item in enumerate(payload[:limit]):
        if not isinstance(item, dict):
            raise ValueError("each slide must be an object")
        title = _clean_text(item.get("title") or item.get("heading") or f"第 {index + 1} 页", 48)
        bullets = _normalize_bullets(item.get("bullets") or item.get("points") or item.get("keywords"))
        image_prompt = _clean_text(item.get("image_prompt") or item.get("prompt"), 1200)
        slides.append(VisualSlide(title=title, bullets=bullets, image_prompt=image_prompt))
    if not slides:
        raise ValueError("slides_json must contain at least one slide")
    return slides


def slides_from_markdown(title: str, outline: str, slide_count: Any = DEFAULT_SLIDE_COUNT) -> list[VisualSlide]:
    content = (outline or "").strip()
    if not content:
        raise ValueError("outline is required when slides_json is empty")
    if len(content) > MAX_OUTLINE_CHARS:
        raise ValueError(f"outline is too long; max {MAX_OUTLINE_CHARS} characters")

    slides: list[VisualSlide] = []
    current: dict[str, Any] | None = None

    for raw_line in content.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        heading = re.match(r"^#{1,3}\s+(.+)$", line)
        if heading:
            if current:
                slides.append(VisualSlide(title=current["title"], bullets=current["bullets"]))
            current = {"title": _clean_text(heading.group(1), 48), "bullets": []}
            continue
        bullet = re.match(r"^[-*+]\s+(.+)$", line)
        if bullet:
            if current is None:
                current = {"title": _clean_text(title or "视觉PPT", 48), "bullets": []}
            if len(current["bullets"]) < 4:
                current["bullets"].append(_clean_text(bullet.group(1), 80))
            continue
        if current is None:
            current = {"title": _clean_text(line, 48), "bullets": []}
        elif len(current["bullets"]) < 4:
            current["bullets"].append(_clean_text(line, 80))

    if current:
        slides.append(VisualSlide(title=current["title"], bullets=current["bullets"]))

    if not slides:
        slides = [VisualSlide(title=_clean_text(title or "视觉PPT", 48), bullets=[_clean_text(content, 80)])]

    return slides[: normalize_slide_count(slide_count)]


STYLE_PRESETS = {
    "mmb_modern_pitch": "modern Chinese startup pitch deck, premium beer technology brand, dark charcoal and warm amber, cinematic lighting, clean geometric composition",
    "beer_brand_campaign": "premium craft beer campaign, lively evening plaza, warm golden light, fresh beer foam, social celebration, high-end commercial poster",
    "finance_pitch": "investor pitch deck, confident business storytelling, elegant data shapes, premium dark blue and gold, refined corporate visual system",
    "tech_saas": "B2B SaaS product strategy deck, clean interface metaphors, devices and cloud infrastructure, crisp modern technology aesthetic",
}


def build_prompt_for_slide(slide: VisualSlide, *, deck_title: str, style_preset: str) -> str:
    style = STYLE_PRESETS.get(style_preset, STYLE_PRESETS["mmb_modern_pitch"])
    bullets = " / ".join(slide.bullets[:4])
    custom = f"\nSpecific visual direction: {slide.image_prompt}" if slide.image_prompt else ""
    return f"""Create a complete 16:9 PowerPoint slide as one polished image.
Deck: {deck_title}
Slide title: {slide.title}
Key message keywords: {bullets or slide.title}
Style: {style}.
Composition: full-bleed presentation slide, strong focal point, clear hierarchy, generous negative space, premium business design.
Text policy: Chinese text is allowed but must be very short, large, and crisp; avoid dense small text, avoid paragraphs, avoid tables of tiny text.
Brand feel: MMB smart fresh beer, self-service beer equipment, energetic but professional.
Output should look like a finished keynote slide, not a poster mockup and not a plain bullet list.{custom}""".strip()


def _image_cache_dir() -> Path:
    root = os.environ.get("VISUAL_PPT_CACHE_DIR") or os.path.join(tempfile.gettempdir(), "visual_ppt_tools_cache")
    path = Path(root)
    path.mkdir(parents=True, exist_ok=True)
    return path


def _cache_key(model: str, quality: str, size: str, prompt: str) -> str:
    payload = json.dumps({"model": model, "quality": quality, "size": size, "prompt": prompt}, ensure_ascii=False, sort_keys=True)
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def _images_endpoint(api_base: str) -> str:
    base = (api_base or "https://api.openai.com/v1").rstrip("/") + "/"
    if base.endswith("/v1/"):
        return urljoin(base, "images/generations")
    return urljoin(base, "v1/images/generations")


def _decode_image_response(response: requests.Response) -> bytes:
    if response.status_code >= 400:
        raise RuntimeError(f"OpenAI image API returned HTTP {response.status_code}: {response.text[:500]}")
    payload = response.json()
    data = payload.get("data") or []
    if not data:
        raise RuntimeError("OpenAI image API returned no data")
    first = data[0]
    if first.get("b64_json"):
        return base64.b64decode(first["b64_json"])
    if first.get("url"):
        image_response = requests.get(first["url"], timeout=120)
        if image_response.status_code >= 400:
            raise RuntimeError(f"OpenAI image URL returned HTTP {image_response.status_code}")
        return image_response.content
    raise RuntimeError("OpenAI image API returned neither b64_json nor url")


def openai_image_client(prompt: str, model: str, quality: str, size: str, config: OpenAIImageConfig) -> bytes:
    if not config.api_key:
        raise RuntimeError("OpenAI API key is required for image generation")

    endpoint = _images_endpoint(config.api_base)
    headers = {"Authorization": f"Bearer {config.api_key}", "Content-Type": "application/json"}
    body = {
        "model": model,
        "prompt": prompt,
        "size": size,
        "quality": quality,
        "n": 1,
        "response_format": "b64_json",
    }
    response = requests.post(endpoint, headers=headers, json=body, timeout=180)
    if response.status_code == 400 and size != FALLBACK_IMAGE_SIZE:
        body["size"] = FALLBACK_IMAGE_SIZE
        response = requests.post(endpoint, headers=headers, json=body, timeout=180)
    return _decode_image_response(response)


def normalize_slide_image(image_bytes: bytes) -> bytes:
    with Image.open(io.BytesIO(image_bytes)) as image:
        image = image.convert("RGB")
        target_w, target_h = TARGET_IMAGE_SIZE
        src_w, src_h = image.size
        scale = max(target_w / src_w, target_h / src_h)
        resized = image.resize((int(src_w * scale), int(src_h * scale)), Image.Resampling.LANCZOS)
        left = max(0, (resized.width - target_w) // 2)
        top = max(0, (resized.height - target_h) // 2)
        cropped = resized.crop((left, top, left + target_w, top + target_h))
        out = io.BytesIO()
        cropped.save(out, format="PNG", optimize=True)
        return out.getvalue()


def get_slide_image(
    *,
    prompt: str,
    config: OpenAIImageConfig,
    image_client: ImageClient | None = None,
) -> bytes:
    client = image_client or (lambda p, m, q, s: openai_image_client(p, m, q, s, config))
    key = _cache_key(config.model, config.quality, config.size, prompt)
    cached = _image_cache_dir() / f"{key}.png"
    if cached.exists():
        return cached.read_bytes()
    raw = client(prompt, config.model, config.quality, config.size)
    normalized = normalize_slide_image(raw)
    cached.write_bytes(normalized)
    return normalized


def normalize_image_concurrency(value: Any, slide_total: int) -> int:
    try:
        configured = int(value or os.environ.get("VISUAL_PPT_IMAGE_CONCURRENCY") or DEFAULT_IMAGE_CONCURRENCY)
    except (TypeError, ValueError):
        configured = DEFAULT_IMAGE_CONCURRENCY
    return max(1, min(slide_total, DEFAULT_IMAGE_CONCURRENCY, configured))


def generate_slide_images(
    *,
    prompts: list[str],
    config: OpenAIImageConfig,
    image_client: ImageClient | None = None,
    concurrency: Any = None,
) -> list[bytes]:
    if not prompts:
        return []

    max_workers = normalize_image_concurrency(concurrency, len(prompts))
    images: list[bytes | None] = [None] * len(prompts)
    with ThreadPoolExecutor(max_workers=max_workers) as executor:
        futures = {
            executor.submit(get_slide_image, prompt=prompt, config=config, image_client=image_client): index
            for index, prompt in enumerate(prompts)
        }
        for future in as_completed(futures):
            images[futures[future]] = future.result()

    return [image for image in images if image is not None]


def build_visual_ppt_artifact(
    *,
    title: str,
    outline: str = "",
    slides_json: str = "",
    filename: str | None = None,
    slide_count: Any = DEFAULT_SLIDE_COUNT,
    style_preset: str = "mmb_modern_pitch",
    image_config: OpenAIImageConfig | None = None,
    image_client: ImageClient | None = None,
    image_concurrency: Any = None,
) -> Artifact:
    limit = normalize_slide_count(slide_count)
    slides = slides_from_json(slides_json, limit) if slides_json.strip() else slides_from_markdown(title, outline, limit)
    if not slides:
        raise ValueError("at least one slide is required")

    config = image_config or OpenAIImageConfig(api_key=os.environ.get("OPENAI_API_KEY", ""))
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]
    prompts = [
        build_prompt_for_slide(slide_data, deck_title=title or "MMB 视觉PPT", style_preset=style_preset)
        for slide_data in slides
    ]
    images = generate_slide_images(
        prompts=prompts,
        config=config,
        image_client=image_client,
        concurrency=image_concurrency,
    )

    for image in images:
        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(io.BytesIO(image), 0, 0, width=prs.slide_width, height=prs.slide_height)

    out = io.BytesIO()
    prs.save(out)
    blob = out.getvalue()
    return Artifact(
        filename=safe_filename(filename, title or "visual_ppt", ".pptx"),
        blob=blob,
        mime_type=PPTX_MIME,
        summary={
            "format": "pptx",
            "mode": "full_slide_images",
            "slides": len(slides),
            "style": style_preset,
            "image_model": config.model,
            "quality": config.quality,
            "size_bytes": len(blob),
            "prompt_count": len(prompts),
            "image_concurrency": normalize_image_concurrency(image_concurrency, len(prompts)),
        },
    )
